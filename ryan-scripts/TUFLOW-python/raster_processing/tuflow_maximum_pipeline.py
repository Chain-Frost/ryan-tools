#!/usr/bin/env python3
"""Build, clip, compress, and report TUFLOW maximum rasters for v06 and v07.

The pipeline discovers result grids under both version folders, creates a
cell-wise maximum with GDAL's VRT ``max`` pixel function, clips that maximum to
the matching ``Model`` polygon, and writes the final compressed GeoTIFF.

Normal GD model outputs are written to an AEP subdirectory. FG03 and the
combined FG01+FG02 model are clipped with FG Outputzones.gpkg and written to an
AEP subdirectory below ``FG`` with ``-clip`` appended.

Incremental rebuilding is based on fingerprints embedded in every output. A
changed, added, or removed source raster, a changed output-zone GeoPackage, a
different compression profile, or a changed processing algorithm causes the
output to be regenerated. Existing legacy outputs without fingerprints are
rebuilt once so subsequent runs can make reliable incremental decisions.

Examples:
    python tuflow_maximum_pipeline.py --dry-run
    python tuflow_maximum_pipeline.py
    python tuflow_maximum_pipeline.py --aep 01.0p 05.0p 10.0p 20.0p 50.0p
    python tuflow_maximum_pipeline.py --version v06 --model GD01 GD02
    python tuflow_maximum_pipeline.py --grid d_HR_Max V_Max
    python tuflow_maximum_pipeline.py --overwrite
    python tuflow_maximum_pipeline.py --profile efficient
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import sys
import tempfile
import time
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Literal

from osgeo import gdal, ogr, osr

from ryan_library.functions.gdal.raster_processing import (
    DEFAULT_OVERVIEW_LEVELS,
    build_external_overviews as shared_build_external_overviews,
    build_vrt,
    geotiff_creation_options,
)
from ryan_library.functions.tuflow.tuflow_result_naming import (
    format_user_template,
    validate_output_filename,
)


# ---------------------------------------------------------------------------
# Project configuration
# ---------------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).absolute().parent
RESULTS_DIR = SCRIPT_DIR.parent
VERSION_NAMES = ("v06", "v07")

MAIN_ZONES = SCRIPT_DIR / "Model Output Zones.gpkg"
FG_OUTPUT_DIR = SCRIPT_DIR / "FG"
FG_ZONES = FG_OUTPUT_DIR / "FG Outputzones.gpkg"
MODEL_FIELD = "Model"

AEP_NAMES = ("01.0p", "05.0p", "10.0p", "20.0p", "50.0p")
DEFAULT_AEPS = AEP_NAMES
GRID_TYPES = ("d_HR_Max", "h_HR_Max", "V_Max", "Z0_Max")
OUTPUT_TEMPLATE = "RTIO_{version}_{model}_EXG_CCNear_{aep}_MAXIMUM_04M_{grid}.tif"
REPORT_TEMPLATE = "MaxRaster_Summary_{stamp}.xlsx"

FG_COMBINED_MODEL = "FG"
FG_COMBINED_EXTENTS = ("FG01", "FG02")
FG_DIRECT_MODELS = ("FG03",)

OVERVIEW_LEVELS = DEFAULT_OVERVIEW_LEVELS
OVERVIEW_RESAMPLING = "NEAREST"

# Increment this only when the numerical processing method changes.
ALGORITHM_ID = "gdal-vrt-max-clip-v1"


# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ZoneCatalog:
    path: Path
    layer_name: str
    field_name: str
    models: dict[str, str]


@dataclass(frozen=True)
class Job:
    version: str
    model: str
    aep: str
    grid_type: str
    sources: tuple[Path, ...]
    output: Path
    zones: ZoneCatalog
    target_aligned: bool = False
    discovery_error: str | None = None


@dataclass
class JobResult:
    version: str
    model: str
    aep: str
    grid_type: str
    output: Path
    source_count: int
    status: str
    message: str
    runtime_seconds: float = 0.0
    output_size_mb: float = 0.0


@dataclass(frozen=True)
class SourceInfo:
    x_resolution: float
    y_resolution: float
    nodata: float
    data_type: int


@dataclass(frozen=True)
class SourceGroup:
    """A named output zone assembled from one or more source extent folders."""

    model: str
    extents: tuple[str, ...]
    zone_set: Literal["main", "fg"] = "main"


DEFAULT_GROUPS: tuple[SourceGroup, ...] = (
    SourceGroup("FG", FG_COMBINED_EXTENTS, "fg"),
    SourceGroup("FG03", ("FG03",), "fg"),
)


def parse_source_group(value: str) -> SourceGroup:
    """Read a group setting: [main|fg:]OUTPUT=SOURCE1,SOURCE2."""
    left, separator, right = value.partition("=")
    if not separator:
        raise ValueError(f"Invalid group {value!r}; expected [main|fg:]OUTPUT=SOURCE1,SOURCE2.")
    prefix, colon, model = left.partition(":")
    if colon:
        if prefix not in {"main", "fg"}:
            raise ValueError(f"Invalid zone set {prefix!r}; expected main or fg.")
        zone_set: Literal["main", "fg"] = "fg" if prefix == "fg" else "main"
    else:
        model = left
        zone_set = "main"
    model = model.strip()
    extents = tuple(part.strip() for part in right.split(","))
    if not model or not extents or any(not part for part in extents):
        raise ValueError(f"Invalid source group {value!r}.")
    if len(set(name.casefold() for name in extents)) != len(extents):
        raise ValueError(f"Duplicate source extents in group {value!r}.")
    return SourceGroup(model=model, extents=extents, zone_set=zone_set)


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------


def load_zone_catalog(path: Path, field_name: str = MODEL_FIELD) -> ZoneCatalog:
    if not path.is_file():
        raise RuntimeError(f"Output-zone GeoPackage not found: {path}")

    dataset = ogr.Open(str(path), 0)
    if dataset is None:
        raise RuntimeError(f"OGR could not open: {path}")
    if dataset.GetLayerCount() != 1:
        raise RuntimeError(
            f"Expected one layer in {path.name}, found {dataset.GetLayerCount()}."
        )

    layer = dataset.GetLayerByIndex(0)
    definition = layer.GetLayerDefn()
    if definition.GetFieldIndex(field_name) < 0:
        raise RuntimeError(
            f"Layer {layer.GetName()!r} in {path.name} has no {field_name!r} field."
        )

    models: dict[str, str] = {}
    for feature in layer:
        value = feature.GetField(field_name)
        geometry = feature.GetGeometryRef()
        if value is None or not str(value).strip():
            raise RuntimeError(f"Feature {feature.GetFID()} in {path.name} has a blank Model.")
        model = str(value).strip()
        key = model.casefold()
        if key in models:
            raise RuntimeError(f"Duplicate Model value {model!r} in {path.name}.")
        if geometry is None or geometry.IsEmpty():
            raise RuntimeError(f"Model {model!r} has empty geometry in {path.name}.")
        models[key] = model

    layer_name = layer.GetName()
    dataset = None
    if not models:
        raise RuntimeError(f"No model polygons found in {path.name}.")
    return ZoneCatalog(path=path.absolute(), layer_name=layer_name, field_name=field_name, models=models)


def collect_inputs(grids_dir: Path, grid_type: str) -> tuple[Path, ...]:
    suffix = f"_{grid_type}.tif".casefold()
    return tuple(
        path
        for path in sorted(grids_dir.iterdir())
        if path.is_file()
        and path.name.casefold().endswith(suffix)
        and "maximum" not in path.name.casefold()
        and not path.stem.casefold().endswith("_src")
    )


def discover_extent_sources(
    version_root: Path,
    aep: str,
    grid_types: tuple[str, ...],
) -> dict[str, tuple[str, Path, dict[str, tuple[Path, ...]]]]:
    discovered: dict[str, tuple[str, Path, dict[str, tuple[Path, ...]]]] = {}
    if not version_root.is_dir():
        return discovered

    for extent_dir in sorted(version_root.iterdir()):
        grids_dir = extent_dir / aep / "grids"
        if not extent_dir.is_dir() or not grids_dir.is_dir():
            continue
        inputs_by_grid = {
            grid_type: collect_inputs(grids_dir, grid_type)
            for grid_type in grid_types
        }
        if any(inputs_by_grid.values()):
            discovered[extent_dir.name.casefold()] = (
                extent_dir.name,
                grids_dir,
                inputs_by_grid,
            )
    return discovered


def output_path(
    version: str,
    model: str,
    aep: str,
    grid_type: str,
    fg: bool,
    *,
    output_root: Path = SCRIPT_DIR,
    template: str = OUTPUT_TEMPLATE,
) -> Path:
    filename = validate_output_filename(
        format_user_template(
            template=template,
            values={"version": version, "model": model, "aep": aep, "grid": grid_type},
            description="TUFLOW maximum output name",
        )
    )
    if fg:
        path = output_root / "FG" / aep / filename
        return path.with_name(f"{path.stem}-clip{path.suffix}")
    return output_root / aep / filename


def discover_jobs(
    versions: tuple[str, ...],
    aeps: tuple[str, ...],
    grid_types: tuple[str, ...],
    wanted_models: set[str] | None,
    main_zones: ZoneCatalog,
    fg_zones: ZoneCatalog | None,
    *,
    input_root: Path = RESULTS_DIR,
    output_root: Path = SCRIPT_DIR,
    output_template: str = OUTPUT_TEMPLATE,
    groups: tuple[SourceGroup, ...] = DEFAULT_GROUPS,
) -> tuple[list[Job], list[str]]:
    """Discover direct model zones and configurable composite output zones.

    A group is a source-folder mapping, not a hard-coded FG algorithm.
    The original FG01/FG02 and FG03 rules remain the defaults.
    """
    jobs: list[Job] = []
    notes: list[str] = []
    group_models: set[str] = set()
    for group in groups:
        key = group.model.casefold()
        if key in group_models:
            raise ValueError(f"Duplicate output group {group.model!r}.")
        group_models.add(key)

    def wanted(model: str) -> bool:
        return wanted_models is None or model.casefold() in wanted_models

    for version in versions:
        version_root = input_root / version
        for aep in aeps:
            extents = discover_extent_sources(version_root, aep, grid_types)
            if not extents:
                notes.append(f"{version} {aep}: no extent folders containing {aep}\\grids")
                continue

            handled: set[str] = set()
            # Direct models using the main output-zone GeoPackage.
            for extent_key, (_extent_name, _grids_dir, inputs_by_grid) in extents.items():
                model = main_zones.models.get(extent_key)
                if model is None or not wanted(model) or model.casefold() in group_models:
                    continue
                handled.add(extent_key)
                for grid_type in grid_types:
                    sources = inputs_by_grid[grid_type]
                    if sources:
                        jobs.append(
                            Job(
                                version=version, model=model, aep=aep,
                                grid_type=grid_type, sources=sources,
                                output=output_path(
                                    version, model, aep, grid_type, fg=False,
                                    output_root=output_root, template=output_template,
                                ),
                                zones=main_zones,
                            )
                        )

            # Any number of input extents can contribute to one output zone.
            for group in groups:
                zone_catalog = fg_zones if group.zone_set == "fg" else main_zones
                if zone_catalog is None:
                    raise ValueError(f"Group {group.model!r} needs an FG zone GeoPackage.")
                model = zone_catalog.models.get(group.model.casefold())
                if model is None:
                    notes.append(f"{version} {aep}: group zone {group.model!r} missing from {zone_catalog.path}")
                    continue
                if not wanted(model):
                    continue
                records = [extents.get(name.casefold()) for name in group.extents]
                present = [record for record in records if record is not None]
                if not present:
                    continue
                handled.update(name.casefold() for name in group.extents if name.casefold() in extents)
                for grid_type in grid_types:
                    sources = tuple(source for record in present for source in record[2][grid_type])
                    if not sources:
                        continue
                    missing = [
                        name for name, record in zip(group.extents, records, strict=True)
                        if record is None or not record[2][grid_type]
                    ]
                    error = (
                        f"Combined {model} {aep} maximum requires "
                        f"{', '.join(group.extents)}; missing {grid_type} sources for {', '.join(missing)}."
                        if missing else None
                    )
                    jobs.append(
                        Job(
                            version=version, model=model, aep=aep,
                            grid_type=grid_type, sources=sources,
                            output=output_path(
                                version, model, aep, grid_type,
                                fg=(group.zone_set == "fg"),
                                output_root=output_root, template=output_template,
                            ),
                            zones=zone_catalog,
                            target_aligned=len(group.extents) > 1,
                            discovery_error=error,
                        )
                    )

            ignored = [record[0] for key, record in extents.items() if key not in handled]
            if ignored:
                notes.append(
                    f"{version} {aep}: ignored extents without a selected output-zone model: "
                    + ", ".join(sorted(ignored))
                )

    jobs.sort(key=lambda job: (
        job.version, job.aep.casefold(), job.model.casefold(), job.grid_type.casefold(),
    ))
    return jobs, notes


# ---------------------------------------------------------------------------
# Incremental state
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# Incremental state
# ---------------------------------------------------------------------------


def file_inventory(paths: tuple[Path, ...]) -> list[dict[str, int | str]]:
    inventory: list[dict[str, int | str]] = []
    for path in sorted(paths, key=lambda item: str(item).casefold()):
        stat = path.stat()
        inventory.append(
            {
                "path": str(path.absolute()),
                "size": stat.st_size,
                "mtime_ns": stat.st_mtime_ns,
            }
        )
    return inventory


def digest(value: object) -> str:
    payload = json.dumps(value, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def source_signature(paths: tuple[Path, ...]) -> str:
    return digest(file_inventory(paths))


def zone_signature(path: Path) -> str:
    stat = path.stat()
    return digest(
        {
            "path": str(path.absolute()),
            "size": stat.st_size,
            "mtime_ns": stat.st_mtime_ns,
        }
    )


def output_metadata(path: Path) -> dict[str, str]:
    dataset = gdal.Open(str(path), gdal.GA_ReadOnly)
    if dataset is None:
        return {}
    metadata = dataset.GetMetadata() or {}
    dataset = None
    return metadata


def overview_path(raster: Path) -> Path:
    return Path(f"{raster}.ovr")


def overviews_are_current(raster: Path) -> bool:
    sidecar = overview_path(raster)
    if not sidecar.is_file() or sidecar.stat().st_mtime_ns < raster.stat().st_mtime_ns:
        return False
    dataset = gdal.Open(str(raster), gdal.GA_ReadOnly)
    if dataset is None:
        return False
    count = dataset.GetRasterBand(1).GetOverviewCount()
    dataset = None
    return count >= len(OVERVIEW_LEVELS)


def determine_action(
    job: Job,
    profile: str,
    build_overviews: bool,
    overwrite: bool,
) -> tuple[str, str, str, str]:
    sources_hash = source_signature(job.sources)
    zones_hash = zone_signature(job.zones.path)

    if overwrite:
        return "rebuild", "--overwrite requested", sources_hash, zones_hash
    if not job.output.is_file():
        return "rebuild", "output does not exist", sources_hash, zones_hash

    metadata = output_metadata(job.output)
    expected = {
        "TUFLOW_MAX_ALGORITHM": ALGORITHM_ID,
        "TUFLOW_SOURCE_SIGNATURE": sources_hash,
        "TUFLOW_ZONE_SIGNATURE": zones_hash,
        "TUFLOW_MODEL_ZONE": job.model,
        "TUFLOW_AEP": job.aep,
        "TUFLOW_STORAGE_PROFILE": profile,
    }
    mismatches = [key for key, value in expected.items() if metadata.get(key) != value]
    if mismatches:
        return (
            "rebuild",
            "missing or changed incremental metadata: " + ", ".join(mismatches),
            sources_hash,
            zones_hash,
        )

    if build_overviews and not overviews_are_current(job.output):
        return "overviews", "raster is current but external overviews are missing or stale", sources_hash, zones_hash
    return "skip", "source rasters, output zones, and processing settings are unchanged", sources_hash, zones_hash


# ---------------------------------------------------------------------------
# Maximum VRT and clipping
# ---------------------------------------------------------------------------


def inspect_sources(paths: tuple[Path, ...]) -> SourceInfo:
    datasets: list[gdal.Dataset] = []
    x_resolutions: list[float] = []
    y_resolutions: list[float] = []
    nodata_values: list[float | None] = []
    data_types: list[int] = []

    try:
        for path in paths:
            dataset = gdal.Open(str(path), gdal.GA_ReadOnly)
            if dataset is None:
                raise RuntimeError(f"GDAL could not open source raster: {path}")
            if dataset.RasterCount != 1:
                raise RuntimeError(f"Expected one band in {path.name}.")
            transform = dataset.GetGeoTransform()
            if transform[2] != 0.0 or transform[4] != 0.0:
                raise RuntimeError(f"Rotated raster grids are not supported: {path.name}")
            if transform[1] <= 0.0 or transform[5] >= 0.0:
                raise RuntimeError(f"Expected a north-up raster: {path.name}")

            band = dataset.GetRasterBand(1)
            datasets.append(dataset)
            x_resolutions.append(abs(transform[1]))
            y_resolutions.append(abs(transform[5]))
            nodata_values.append(band.GetNoDataValue())
            data_types.append(band.DataType)

        first_srs = datasets[0].GetSpatialRef()
        if first_srs is None:
            raise RuntimeError(f"Source raster has no CRS: {paths[0].name}")
        for path, dataset in zip(paths[1:], datasets[1:]):
            spatial_ref = dataset.GetSpatialRef()
            if spatial_ref is None or not first_srs.IsSame(spatial_ref):
                raise RuntimeError(f"Source raster CRS differs: {path.name}")

        if any(value is None for value in nodata_values):
            raise RuntimeError("Every source raster must define a NoData value.")
        nodata = float(nodata_values[0])
        if not all(math.isclose(float(value), nodata) for value in nodata_values[1:]):
            raise RuntimeError(f"Source NoData values differ: {nodata_values}")
        if len(set(data_types)) != 1:
            names = [gdal.GetDataTypeName(value) for value in data_types]
            raise RuntimeError(f"Source raster data types differ: {names}")

        return SourceInfo(
            x_resolution=min(x_resolutions),
            y_resolution=min(y_resolutions),
            nodata=nodata,
            data_type=data_types[0],
        )
    finally:
        datasets.clear()


def create_maximum_vrt(job: Job, source_info: SourceInfo) -> tuple[gdal.Dataset, str]:
    """Build a GDAL 3.12+ maximum VRT through the shared library helper."""
    virtual_path = f"/vsimem/tuflow_max_{uuid.uuid4().hex}.vrt"
    try:
        if job.target_aligned:
            build_vrt(
                job.sources,
                Path(virtual_path),
                source_nodata=source_info.nodata,
                vrt_nodata=source_info.nodata,
                pixel_function="max",
                resample_alg=gdal.GRA_NearestNeighbour,
                strict=True,
                resolution="user",
                x_res=source_info.x_resolution,
                y_res=source_info.y_resolution,
                target_aligned_pixels=True,
            )
        else:
            build_vrt(
                job.sources,
                Path(virtual_path),
                source_nodata=source_info.nodata,
                vrt_nodata=source_info.nodata,
                pixel_function="max",
                resample_alg=gdal.GRA_NearestNeighbour,
                strict=True,
                resolution="highest",
            )

        # The shared helper closes its output. Keep the reopened VRT alive
        # until the polygon clipping has materialised.
        maximum = gdal.Open(virtual_path, gdal.GA_ReadOnly)
        if maximum is None:
            raise RuntimeError("GDAL failed to reopen the maximum VRT.")
        return maximum, virtual_path
    except Exception:
        gdal.Unlink(virtual_path)
        raise


def temporary_tif(destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(
        prefix=f".{destination.stem}_building_",
        suffix=destination.suffix,
        dir=destination.parent,
        delete=False,
    )
    path = Path(handle.name)
    handle.close()
    path.unlink()
    return path


def selected_zone_geometry(job: Job) -> tuple[ogr.DataSource, ogr.Layer, ogr.Geometry]:
    vector_dataset = ogr.Open(str(job.zones.path), 0)
    if vector_dataset is None:
        raise RuntimeError(f"OGR could not open output zones: {job.zones.path}")
    layer = vector_dataset.GetLayerByName(job.zones.layer_name)
    if layer is None:
        raise RuntimeError(f"Layer {job.zones.layer_name!r} was not found.")
    escaped = job.model.replace("'", "''")
    layer.SetAttributeFilter(f'"{job.zones.field_name}" = \'{escaped}\'')
    feature = layer.GetNextFeature()
    if feature is None:
        raise RuntimeError(f"No polygon found for Model={job.model!r}.")
    geometry = feature.GetGeometryRef()
    if geometry is None or geometry.IsEmpty():
        raise RuntimeError(f"Polygon for Model={job.model!r} is empty.")
    return vector_dataset, layer, geometry.Clone()


def clip_bounds_and_crs(
    maximum: gdal.Dataset,
    job: Job,
) -> tuple[tuple[float, float, float, float], str | None]:
    transform = maximum.GetGeoTransform()
    if transform[2] != 0.0 or transform[4] != 0.0:
        raise RuntimeError("Rotated maximum grids are not supported.")

    vector_dataset, layer, geometry = selected_zone_geometry(job)
    source_wkt = maximum.GetProjection()
    source_srs = osr.SpatialReference()
    if source_wkt:
        source_srs.ImportFromWkt(source_wkt)
    zone_srs = layer.GetSpatialRef()
    if not source_wkt or zone_srs is None:
        raise RuntimeError("The maximum raster and output zones must both define a CRS.")

    source_srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    zone_srs = zone_srs.Clone()
    zone_srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)

    source_name = (source_srs.GetName() or "").casefold()
    legacy_gda94_mg = bool(source_srs.IsLocal() and source_name == "gda94_mg")
    source_srs_override: str | None = None
    if legacy_gda94_mg:
        # TUFLOW/ASC_to_ASC rasters use correct MGA50 coordinates but sometimes
        # carry LOCAL_CS["GDA94_MG"]. Use the polygon's EPSG definition.
        source_srs_override = zone_srs.ExportToWkt()
    elif source_srs.IsLocal():
        raise RuntimeError(
            f"Unsupported local raster CRS {source_srs.GetName()!r}; expected GDA94_MG."
        )
    elif not zone_srs.IsSame(source_srs):
        geometry.Transform(osr.CoordinateTransformation(zone_srs, source_srs))

    min_x, max_x, min_y, max_y = geometry.GetEnvelope()
    origin_x, pixel_x, _, origin_y, _, pixel_y = transform
    pixel_y_abs = abs(pixel_y)
    first_column = max(0, math.floor((min_x - origin_x) / pixel_x))
    last_column = min(
        maximum.RasterXSize,
        math.ceil((max_x - origin_x) / pixel_x),
    )
    first_row = max(0, math.floor((origin_y - max_y) / pixel_y_abs))
    last_row = min(
        maximum.RasterYSize,
        math.ceil((origin_y - min_y) / pixel_y_abs),
    )
    if first_column >= last_column or first_row >= last_row:
        raise RuntimeError(f"Model {job.model!r} polygon does not overlap its maximum raster.")

    bounds = (
        origin_x + first_column * pixel_x,
        origin_y - last_row * pixel_y_abs,
        origin_x + last_column * pixel_x,
        origin_y - first_row * pixel_y_abs,
    )
    vector_dataset = None
    return bounds, source_srs_override


def build_external_overviews(raster: Path, profile: str) -> None:
    """Retry the shared GDAL overview helper for slow network file shares."""
    sidecar = overview_path(raster)
    sidecar_existed = sidecar.is_file()
    failures: list[str] = []
    for attempt in range(1, 4):
        if attempt > 1:
            time.sleep(float(attempt - 1))
        try:
            shared_build_external_overviews(
                raster,
                levels=OVERVIEW_LEVELS,
                resampling="nearest",
                profile=profile,
                refresh=True,
            )
            return
        except (RuntimeError, OSError) as error:
            failures.append(f"attempt {attempt}: {error}")
            if not sidecar_existed:
                sidecar.unlink(missing_ok=True)
    raise RuntimeError(
        "GDAL failed to build external overviews after 3 attempts: " + "; ".join(failures)
    )

def materialize_clipped_maximum(
    maximum: gdal.Dataset,
    source_info: SourceInfo,
    job: Job,
    destination: Path,
    profile: str,
    sources_hash: str,
    zones_hash: str,
) -> None:
    bounds, source_srs_override = clip_bounds_and_crs(maximum, job)
    transform = maximum.GetGeoTransform()
    escaped = job.model.replace("'", "''")

    warp_options: dict[str, object] = {
        "format": "GTiff",
        "outputBounds": bounds,
        "xRes": abs(transform[1]),
        "yRes": abs(transform[5]),
        "resampleAlg": gdal.GRA_NearestNeighbour,
        "cutlineDSName": str(job.zones.path),
        "cutlineLayer": job.zones.layer_name,
        "cutlineWhere": f'"{job.zones.field_name}" = \'{escaped}\'',
        "cropToCutline": False,
        "srcNodata": source_info.nodata,
        "dstNodata": source_info.nodata,
        "multithread": True,
        "creationOptions": geotiff_creation_options(profile, threads="ALL_CPUS"),
        "warpOptions": ["NUM_THREADS=ALL_CPUS", "UNIFIED_SRC_NODATA=YES"],
    }
    if source_srs_override is not None:
        warp_options.update(srcSRS=source_srs_override, dstSRS=source_srs_override)

    output = gdal.Warp(
        str(destination),
        maximum,
        options=gdal.WarpOptions(**warp_options),
    )
    if output is None:
        raise RuntimeError(gdal.GetLastErrorMsg() or "GDAL did not create clipped output.")

    output.SetMetadataItem("TUFLOW_MAX_ALGORITHM", ALGORITHM_ID)
    output.SetMetadataItem("TUFLOW_SOURCE_SIGNATURE", sources_hash)
    output.SetMetadataItem("TUFLOW_SOURCE_COUNT", str(len(job.sources)))
    output.SetMetadataItem("TUFLOW_ZONE_SIGNATURE", zones_hash)
    output.SetMetadataItem("TUFLOW_MODEL_ZONE", job.model)
    output.SetMetadataItem("TUFLOW_MODEL_ZONES_SOURCE", str(job.zones.path))
    output.SetMetadataItem("TUFLOW_STORAGE_PROFILE", profile)
    output.SetMetadataItem("TUFLOW_VERSION", job.version)
    output.SetMetadataItem("TUFLOW_AEP", job.aep)
    output.SetMetadataItem("TUFLOW_GRID_TYPE", job.grid_type)
    output.SetMetadataItem("TUFLOW_MAXIMUM_METHOD", "GDAL VRT pixel function max")
    output.SetMetadataItem("TUFLOW_RESAMPLING", "nearest neighbour")
    if source_srs_override is not None:
        output.SetMetadataItem("TUFLOW_SOURCE_CRS_OVERRIDE", "GDA94_MG -> output-zone CRS")
    output.FlushCache()
    output = None


def publish_output(temporary: Path, destination: Path, build_overviews: bool) -> None:
    """Install a raster/overview pair, restoring prior files on failure.

    Two independent filesystem renames are not fully atomic. Backing up the
    existing pair before publishing lets an ordinary exception restore the
    previous consistent generation rather than leave a new TIFF with an old
    overview. Hard power loss still needs manual recovery.
    """
    temporary_overview = overview_path(temporary)
    destination_overview = overview_path(destination)
    if not temporary.is_file():
        raise RuntimeError(f"Temporary raster is missing: {temporary}")
    if build_overviews and not temporary_overview.is_file():
        raise RuntimeError(f"Temporary overview is missing: {temporary_overview}")

    marker = uuid.uuid4().hex
    raster_backup = destination.with_name(f".{destination.stem}.{marker}.bak{destination.suffix}")
    overview_backup = Path(f"{raster_backup}.ovr")
    backed_up_raster = False
    backed_up_overview = False
    try:
        if destination.is_file():
            os.replace(destination, raster_backup)
            backed_up_raster = True
        if destination_overview.is_file():
            os.replace(destination_overview, overview_backup)
            backed_up_overview = True
        os.replace(temporary, destination)
        if build_overviews:
            os.replace(temporary_overview, destination_overview)
    except Exception:
        # Restore both files as a pair. If a restoration itself fails, preserve
        # the backup paths rather than deleting valuable prior results.
        destination.unlink(missing_ok=True)
        destination_overview.unlink(missing_ok=True)
        if backed_up_raster:
            os.replace(raster_backup, destination)
        if backed_up_overview:
            os.replace(overview_backup, destination_overview)
        raise
    else:
        raster_backup.unlink(missing_ok=True)
        overview_backup.unlink(missing_ok=True)


def run_build(
    job: Job,
    profile: str,
    build_overviews: bool,
    sources_hash: str,
    zones_hash: str,
) -> tuple[float, float]:
    started = time.perf_counter()
    temporary = temporary_tif(job.output)
    virtual_path = ""
    maximum: gdal.Dataset | None = None
    try:
        source_info = inspect_sources(job.sources)
        maximum, virtual_path = create_maximum_vrt(job, source_info)
        materialize_clipped_maximum(
            maximum,
            source_info,
            job,
            temporary,
            profile,
            sources_hash,
            zones_hash,
        )
        maximum = None
        gdal.Unlink(virtual_path)
        virtual_path = ""
        if build_overviews:
            build_external_overviews(temporary, profile)
        publish_output(temporary, job.output, build_overviews)
        elapsed = time.perf_counter() - started
        return elapsed, job.output.stat().st_size / 1e6
    finally:
        maximum = None
        if virtual_path:
            gdal.Unlink(virtual_path)
        temporary.unlink(missing_ok=True)
        overview_path(temporary).unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------


def write_report(
    path: Path,
    results: list[JobResult],
    jobs: list[Job],
    run_info: list[tuple[str, str]],
) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("! openpyxl is not installed; the Excel summary was skipped.")
        return False

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    status_fills = {
        "BUILT": PatternFill("solid", fgColor="E2EFDA"),
        "OVERVIEWS": PatternFill("solid", fgColor="FFF2CC"),
        "SKIPPED": PatternFill("solid", fgColor="F2F2F2"),
        "FAILED": PatternFill("solid", fgColor="FCE4EC"),
        "DRY-RUN": PatternFill("solid", fgColor="DDEBF7"),
    }

    def table(worksheet, headers, rows, widths, status_column=None):
        worksheet.append(headers)
        for index, _header in enumerate(headers, start=1):
            cell = worksheet.cell(1, index)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for row in rows:
            worksheet.append(row)
        for row_index in range(2, worksheet.max_row + 1):
            status = (
                worksheet.cell(row_index, status_column).value
                if status_column is not None
                else None
            )
            for column_index in range(1, len(headers) + 1):
                cell = worksheet.cell(row_index, column_index)
                cell.font = body_font
                if status in status_fills:
                    cell.fill = status_fills[status]
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width
        worksheet.freeze_panes = "A2"
        if worksheet.max_row >= 2:
            worksheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
            )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    grouped: dict[tuple[str, str, str], dict[str, int]] = {}
    for result in results:
        counts = grouped.setdefault((result.version, result.aep, result.model), {})
        counts[result.status] = counts.get(result.status, 0) + 1
    summary_rows = [
        [
            version,
            aep,
            model,
            sum(counts.values()),
            counts.get("BUILT", 0),
            counts.get("OVERVIEWS", 0),
            counts.get("SKIPPED", 0),
            counts.get("DRY-RUN", 0),
            counts.get("FAILED", 0),
        ]
        for (version, aep, model), counts in sorted(grouped.items())
    ]
    table(
        summary_sheet,
        [
            "Version",
            "AEP",
            "Model",
            "Jobs",
            "Built",
            "Overviews",
            "Skipped",
            "Dry run",
            "Failed",
        ],
        summary_rows,
        [12, 12, 20, 10, 10, 12, 10, 10, 10],
    )

    jobs_sheet = workbook.create_sheet("Jobs")
    job_rows = [
        [
            result.version,
            result.aep,
            result.model,
            result.grid_type,
            result.output.name,
            result.source_count,
            result.status,
            result.message,
            round(result.runtime_seconds, 1) if result.runtime_seconds else None,
            round(result.output_size_mb, 2) if result.output_size_mb else None,
            str(result.output.parent),
        ]
        for result in results
    ]
    table(
        jobs_sheet,
        [
            "Version",
            "AEP",
            "Model",
            "Grid type",
            "Output raster",
            "Source count",
            "Status",
            "Message",
            "Runtime (s)",
            "Output size (MB)",
            "Output folder",
        ],
        job_rows,
        [12, 12, 20, 14, 72, 14, 12, 68, 14, 18, 72],
        status_column=7,
    )

    inputs_sheet = workbook.create_sheet("Input Rasters")
    input_rows = []
    for job in jobs:
        for source in job.sources:
            stat = source.stat()
            input_rows.append(
                [
                    job.version,
                    job.aep,
                    job.model,
                    job.grid_type,
                    source.name,
                    round(stat.st_size / 1e6, 2),
                    datetime.fromtimestamp(stat.st_mtime),
                    str(source.parent),
                    job.output.name,
                ]
            )
    table(
        inputs_sheet,
        [
            "Version",
            "AEP",
            "Model",
            "Grid type",
            "Input raster",
            "Size (MB)",
            "Last modified",
            "Input folder",
            "Feeds into output",
        ],
        input_rows,
        [12, 12, 20, 14, 76, 12, 20, 80, 72],
    )
    for row_index in range(2, inputs_sheet.max_row + 1):
        inputs_sheet.cell(row_index, 7).number_format = "yyyy-mm-dd hh:mm:ss"

    info_sheet = workbook.create_sheet("Run Info")
    info_sheet["A1"] = "TUFLOW maximum raster pipeline"
    info_sheet["A1"].font = Font(name="Arial", size=12, bold=True)
    for row_index, (label, value) in enumerate(run_info, start=3):
        info_sheet.cell(row_index, 1, label).font = Font(name="Arial", size=10, bold=True)
        info_sheet.cell(row_index, 2, value).font = body_font
    info_sheet.column_dimensions["A"].width = 30
    info_sheet.column_dimensions["B"].width = 110

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return True


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "--version",
        nargs="+",
        help="Only search selected version folders (default: v06 v07).",
    )
    parser.add_argument(
        "--aep",
        nargs="+",
        help="Only process selected AEP folders (default: all configured AEPs).",
    )
    parser.add_argument(
        "--model",
        "--extent",
        dest="models",
        nargs="+",
        metavar="NAME",
        help="Only process selected output models, for example GD01 FG FG03.",
    )
    parser.add_argument(
        "--grid",
        nargs="+",
        help="Only process selected grid types.",
    )
    parser.add_argument(
        "--profile",
        choices=("tuflow", "efficient"),
        default="tuflow",
        help="Compression profile: conservative DEFLATE or tiled ZSTD (default: tuflow).",
    )
    parser.add_argument("--input-root", type=Path, default=RESULTS_DIR, help="Folder containing version subdirectories.")
    parser.add_argument("--output-root", type=Path, default=SCRIPT_DIR, help="Output root (FG results under its FG subfolder).")
    parser.add_argument("--main-zones", type=Path, default=MAIN_ZONES, help="Main model output-zone GeoPackage (one layer).")
    parser.add_argument("--fg-zones", type=Path, default=FG_ZONES, help="Special output-zone GeoPackage (one layer).")
    parser.add_argument("--zone-field", default=MODEL_FIELD, help="Polygon attribute used to match model names.")
    parser.add_argument("--output-template", default=OUTPUT_TEMPLATE, help="Filename template with version/model/aep/grid.")
    parser.add_argument(
        "--group", action="append", metavar="[main|fg:]MODEL=EXTENT1,EXTENT2",
        help="Source extents contributing to an output zone. When supplied, replaces default FG groups.",
    )
    parser.add_argument("--overwrite", action="store_true", help="Rebuild every discovered output.")
    parser.add_argument("--dry-run", action="store_true", help="Report intended actions without writing rasters.")
    parser.add_argument("--no-overviews", action="store_true", help="Do not create external .ovr pyramids.")
    parser.add_argument("--no-report", action="store_true", help="Do not create the Excel summary.")
    parser.add_argument("--report", type=Path, help="Write the Excel summary to this path.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    gdal.UseExceptions()
    ogr.UseExceptions()

    run_started = datetime.now()
    versions = tuple(args.version or VERSION_NAMES)
    aeps = tuple(dict.fromkeys(args.aep or DEFAULT_AEPS))
    grid_types = tuple(args.grid or GRID_TYPES)
    wanted_models = (
        {value.casefold() for value in args.models}
        if args.models
        else None
    )
    build_overviews = not args.no_overviews

    try:
        groups = tuple(parse_source_group(value) for value in args.group) if args.group is not None else DEFAULT_GROUPS
        main_zones = load_zone_catalog(args.main_zones.absolute(), args.zone_field)
        fg_zones = (
            load_zone_catalog(args.fg_zones.absolute(), args.zone_field)
            if any(group.zone_set == "fg" for group in groups) else None
        )
        jobs, discovery_notes = discover_jobs(
            versions,
            aeps,
            grid_types,
            wanted_models,
            main_zones,
            fg_zones,
            input_root=args.input_root.absolute(),
            output_root=args.output_root.absolute(),
            output_template=args.output_template,
            groups=groups,
        )
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print(f"Versions : {', '.join(versions)}")
    print(f"AEPs     : {', '.join(aeps)}")
    print(f"Grid types: {', '.join(grid_types)}")
    print(f"Profile  : {args.profile}")
    print(f"Overviews: {'enabled' if build_overviews else 'disabled'}")
    print(f"Main zones: {args.main_zones.absolute()}")
    print(f"FG zones  : {args.fg_zones.absolute() if fg_zones is not None else 'not used'}")
    print(f"Discovered {len(jobs)} output job(s).")
    for note in discovery_notes:
        print(f"  NOTE: {note}")
    if not jobs:
        print("ERROR: No matching maximum-raster jobs were discovered.", file=sys.stderr)
        return 1

    results: list[JobResult] = []
    for index, job in enumerate(jobs, start=1):
        label = (
            f"[{index}/{len(jobs)}] {job.version} {job.aep} "
            f"{job.model} {job.grid_type}"
        )
        print(f"\n{label}")
        print(f"  Sources: {len(job.sources)}")
        print(f"  Output : {job.output}")

        if job.discovery_error:
            print(f"  FAILED: {job.discovery_error}")
            results.append(
                JobResult(
                    job.version,
                    job.model,
                    job.aep,
                    job.grid_type,
                    job.output,
                    len(job.sources),
                    "FAILED",
                    job.discovery_error,
                )
            )
            continue

        try:
            action, reason, sources_hash, zones_hash = determine_action(
                job,
                args.profile,
                build_overviews,
                args.overwrite,
            )
            print(f"  Action : {action} ({reason})")

            if args.dry_run:
                results.append(
                    JobResult(
                        job.version,
                        job.model,
                        job.aep,
                        job.grid_type,
                        job.output,
                        len(job.sources),
                        "DRY-RUN",
                        f"Would {action}: {reason}",
                        output_size_mb=(job.output.stat().st_size / 1e6 if job.output.is_file() else 0.0),
                    )
                )
                continue

            if action == "skip":
                results.append(
                    JobResult(
                        job.version,
                        job.model,
                        job.aep,
                        job.grid_type,
                        job.output,
                        len(job.sources),
                        "SKIPPED",
                        reason,
                        output_size_mb=job.output.stat().st_size / 1e6,
                    )
                )
                continue

            if action == "overviews":
                started = time.perf_counter()
                build_external_overviews(job.output, args.profile)
                elapsed = time.perf_counter() - started
                results.append(
                    JobResult(
                        job.version,
                        job.model,
                        job.aep,
                        job.grid_type,
                        job.output,
                        len(job.sources),
                        "OVERVIEWS",
                        reason,
                        elapsed,
                        job.output.stat().st_size / 1e6,
                    )
                )
                print(f"  OK: external overviews refreshed in {elapsed:.1f}s")
                continue

            elapsed, size_mb = run_build(
                job,
                args.profile,
                build_overviews,
                sources_hash,
                zones_hash,
            )
            results.append(
                JobResult(
                    job.version,
                    job.model,
                    job.aep,
                    job.grid_type,
                    job.output,
                    len(job.sources),
                    "BUILT",
                    reason,
                    elapsed,
                    size_mb,
                )
            )
            print(f"  OK: built {size_mb:.1f} MB in {elapsed:.1f}s")
        except Exception as error:
            print(f"  FAILED: {error}", file=sys.stderr)
            results.append(
                JobResult(
                    job.version,
                    job.model,
                    job.aep,
                    job.grid_type,
                    job.output,
                    len(job.sources),
                    "FAILED",
                    str(error),
                )
            )

    counts: dict[str, int] = {}
    for result in results:
        counts[result.status] = counts.get(result.status, 0) + 1
    print("\nFinished: " + ", ".join(f"{status}={count}" for status, count in sorted(counts.items())))

    if not args.no_report:
        report_path = args.report or args.output_root.absolute() / REPORT_TEMPLATE.format(
            stamp=run_started.strftime("%Y%m%d_%H%M%S")
        )
        run_info = [
            ("Run started", run_started.strftime("%Y-%m-%d %H:%M:%S")),
            ("Run finished", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Mode", "Dry run" if args.dry_run else "Live"),
            ("Algorithm", ALGORITHM_ID),
            ("Versions searched", ", ".join(versions)),
            ("AEPs searched", ", ".join(aeps)),
            ("Grid types", ", ".join(grid_types)),
            ("Selected models", ", ".join(args.models) if args.models else "All discovered"),
            ("Compression profile", args.profile),
            ("External overview levels", ", ".join(map(str, OVERVIEW_LEVELS)) if build_overviews else "Disabled"),
            ("Main output root", str(args.output_root.absolute() / "<AEP>")),
            ("FG output root", str(args.output_root.absolute() / "FG" / "<AEP>")),
            ("Main model zones", str(args.main_zones.absolute())),
            ("FG model zones", str(args.fg_zones.absolute()) if fg_zones is not None else "not used"),
            ("Overwrite", "Yes" if args.overwrite else "No"),
            ("Jobs discovered", str(len(jobs))),
            ("Results", ", ".join(f"{key}={value}" for key, value in sorted(counts.items()))),
        ]
        try:
            if write_report(report_path, results, jobs, run_info):
                print(f"Summary written to: {report_path}")
        except PermissionError:
            print(f"! Could not write summary; is it open in Excel? {report_path}")
        except Exception as error:
            print(f"! Could not write summary: {error}")

    return 1 if counts.get("FAILED", 0) else 0


if __name__ == "__main__":
    raise SystemExit(main())
