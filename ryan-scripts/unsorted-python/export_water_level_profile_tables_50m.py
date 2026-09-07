r"""Export profile ground and water levels at exact 50 m chainages.

Edit the example input paths, scenarios, and sampling settings below before
running. This standalone script does not require project-specific companion
scripts. Profiles are numbered in scenario and input-feature order. It writes ground level, water level, and
water depth to one ordered CSV per profile, a combined Excel workbook, and a
Word-ready document containing one captioned table per Profile + Scenario.

Run from any directory::

    py -3.14 export_water_level_profile_tables_50m.py
"""

from __future__ import annotations

# GeoPandas, NumPy, Rasterio, and Shapely expose incomplete third-party typing.
# pyright: reportUnknownArgumentType=false, reportUnknownMemberType=false, reportUnknownVariableType=false

from collections.abc import Mapping
from pathlib import Path
from typing import Any, Literal, cast

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.styles.style import ParagraphStyle
from docx.styles.style import _TableStyle  # pyright: ignore[reportPrivateUsage]
from loguru import logger
import numpy as np
from numpy.typing import NDArray
import pandas as pd
from shapely.geometry import LineString
from shapely.geometry.base import BaseGeometry

from ryan_library.functions.gdal.profiling import (
    RasterSamplingError,
    interpolate_short_nan_gaps,
    sample_raster_along_line,
)
from ryan_library.functions.loguru_helpers import configure_serial_logging
from ryan_library.functions.path_stuff import sanitize_windows_filename
from ryan_library.orchestrators.tuflow import water_level_profiles as profiles

SCRIPT_VERSION = "2026-09-02.3"

# Example inputs: replace these with your own model paths and scenario names.
INPUT_DIRECTORY = Path("C:/path/to/model")
SCENARIOS: dict[str, tuple[Path, Path]] = {
    "Scenario A": (INPUT_DIRECTORY / "profiles.gpkg", INPUT_DIRECTORY / "results" / "scenario_a"),
}
TERRAIN_TIF = INPUT_DIRECTORY / "terrain.tif"
LINES_LAYER_NAME: str | None = None
NAME_FIELD = "Name"
LINES_CRS_IF_MISSING: str | None = None
TARGET_AEPS = ("01.00p", "20.00p")
TARGET_RESULT_TYPE = "h_HR_Max"
SPACING: float | None = None
SAMPLING_METHOD: Literal["bilinear", "bilinear_valid", "bilinear_masked", "nearest"] = "bilinear_masked"
MAX_INTERPOLATION_GAP = 10
DRY_AREA_HANDLING: Literal["no_plot", "ground_level"] = "no_plot"
DISCONNECTED_LINE_HANDLING: Literal["error", "separate"] = "error"

SECTION_INTERVAL_M = 50.0
OUTPUT_DIRECTORY = Path(__file__).resolve().parent / "profile_tables_50m"
WORKBOOK_NAME = "water_level_profiles_50m.xlsx"
WORD_DOCUMENT_NAME = "water_level_profile_tables_50m_word_appendix_style.docx"
APPENDIX_TABLE_STYLE = "appendix table style"
SUMMARY_CSV_NAME = "water_level_profile_decision_summary.csv"
SUMMARY_PIVOT_CSV_NAME = "water_level_profile_decision_summary_pivot.csv"
DECIMAL_PLACES = 3
WORD_DECIMAL_PLACES = 2
REPORTING_AEP = "01.00p"
SUMMARY_AEPS = ("20.00p", "01.00p")

type FloatArray = NDArray[np.float64]

WORD_TABLE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("Chainage (m)", "Chainage (m)"),
    ("Ground Level (mAHD)", "Ground (mAHD)"),
    ("1% AEP WL (mAHD)", "1% WSEL (mAHD)"),
    ("1% AEP Depth (m)", "1% Depth (m)"),
    ("20% AEP WL (mAHD)", "20% WSEL (mAHD)"),
    ("20% AEP Depth (m)", "20% Depth (m)"),
)


def _set_cell_shading(cell: Any, fill: str) -> None:
    """Apply a solid fill to a Word table cell."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_margins(cell: Any, *, top: int, bottom: int) -> None:
    """Set vertical cell margins in twentieths of a point."""
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for edge, value in (("top", top), ("bottom", bottom)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row: Any) -> None:
    """Repeat a table's header row when it spans Word pages."""
    row_properties = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    row_properties.append(table_header)


def _prevent_row_split(row: Any) -> None:
    """Keep one table row together across a Word page boundary."""
    row_properties = row._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:cantSplit"))


def _set_minimum_row_height(row: Any) -> None:
    """Use Word's minimum content-driven row height setting."""
    row_properties = row._tr.get_or_add_trPr()
    for existing_height in row_properties.findall(qn("w:trHeight")):
        row_properties.remove(existing_height)
    row_height = OxmlElement("w:trHeight")
    row_height.set(qn("w:val"), "0")
    row_height.set(qn("w:hRule"), "atLeast")
    row_properties.append(row_height)


def _configure_word_document(document: DocumentObject) -> None:
    """Set the report page geometry and base typography."""
    section = document.sections[0]
    section.page_width = Cm(21.03)
    section.page_height = Cm(29.70)
    section.top_margin = Cm(2.40)
    section.right_margin = Cm(2.00)
    section.bottom_margin = Cm(1.50)
    section.left_margin = Cm(2.40)
    normal_style = cast(ParagraphStyle, document.styles["Normal"])
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)
    appendix_style = cast(
        _TableStyle,
        document.styles.add_style(
            APPENDIX_TABLE_STYLE,
            WD_STYLE_TYPE.TABLE,
        ),
    )
    appendix_style.base_style = document.styles["Table Grid"]


def _write_word_tables(
    document_path: Path,
    tables: tuple[tuple[int, str, pd.DataFrame], ...],
) -> None:
    """Write one report-formatted table per Profile + Scenario combination."""
    document = Document()
    _configure_word_document(document)

    for table_index, (order, display_name, dataframe) in enumerate(tables):
        scenarios = tuple(dataframe["Scenario"].drop_duplicates())
        profile_names = tuple(dataframe["Profile"].drop_duplicates())
        if len(scenarios) != 1 or len(profile_names) != 1:
            raise ValueError(f"{display_name!r} is not one Profile + Scenario combination.")
        scenario = str(scenarios[0])
        profile_name = str(profile_names[0])
        if table_index:
            document.add_page_break()

        caption = document.add_paragraph()
        caption.paragraph_format.keep_with_next = True
        caption.paragraph_format.space_before = Pt(12)
        caption.paragraph_format.space_after = Pt(4)
        caption_run = caption.add_run(
            f"Table {order}: Water Level Profile Data - " f"Profile: {profile_name}; Scenario: {scenario}"
        )
        caption_run.bold = True

        word_table = document.add_table(rows=1, cols=len(WORD_TABLE_COLUMNS))
        word_table.style = APPENDIX_TABLE_STYLE
        word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        word_table.autofit = True
        header_row = word_table.rows[0]
        _set_repeat_table_header(header_row)
        _prevent_row_split(header_row)
        _set_minimum_row_height(header_row)
        for cell, (_, report_header) in zip(header_row.cells, WORD_TABLE_COLUMNS, strict=True):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_shading(cell, "DBDBDB")
            _set_cell_margins(cell, top=60, bottom=60)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(report_header)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)

        for _, source_row in dataframe.iterrows():
            word_row = word_table.add_row()
            _prevent_row_split(word_row)
            _set_minimum_row_height(word_row)
            for cell, (source_header, _) in zip(word_row.cells, WORD_TABLE_COLUMNS, strict=True):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_cell_margins(cell, top=60, bottom=60)
                value = source_row[source_header]
                if pd.isna(value):
                    text = ""
                elif source_header == "Chainage (m)":
                    text = f"{float(value):.0f}"
                else:
                    text = f"{float(value):.{WORD_DECIMAL_PLACES}f}"
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = paragraph.add_run(text)
                run.font.name = "Calibri"
                run.font.size = Pt(10)

    document.save(str(document_path))


def _aep_column_name(aep: str) -> str:
    """Return a readable water-level column heading for a TUFLOW AEP token."""
    numeric = aep.casefold().removesuffix("p")
    try:
        label = f"{float(numeric):g}% AEP"
    except ValueError:
        label = f"{aep} AEP"
    return f"{label} WL (mAHD)"


def _depth_column_name(aep: str) -> str:
    """Return a readable water-depth column heading for a TUFLOW AEP token."""
    return _aep_column_name(aep).replace("WL (mAHD)", "Depth (m)")


def _event_label(aep: str) -> str:
    """Return an AEP label with its annual-recurrence-interval equivalent."""
    numeric = aep.casefold().removesuffix("p")
    try:
        percentage = float(numeric)
    except ValueError:
        return f"{aep} AEP"
    if percentage <= 0.0:
        raise ValueError(f"AEP must be positive, received {aep!r}.")
    return f"{percentage:g}% AEP (1 in {100.0 / percentage:g})"


def _exact_chainages(line_length: float) -> FloatArray:
    """Return 0, 50, 100, ... stations that do not exceed the line length."""
    if not np.isfinite(line_length) or line_length <= 0.0:
        raise ValueError(f"Profile length must be positive and finite, received {line_length!r}.")
    final_section = int(np.floor((line_length + 1e-9) / SECTION_INTERVAL_M))
    return np.arange(final_section + 1, dtype=np.float64) * SECTION_INTERVAL_M


def _interpolate_at_chainages(
    source_chainages: FloatArray,
    source_values: FloatArray,
    target_chainages: FloatArray,
) -> FloatArray:
    """Interpolate without bridging across NoData values."""
    if source_chainages.shape != source_values.shape:
        raise ValueError("Source chainage and value arrays must have matching shapes.")
    result = np.full(target_chainages.shape, np.nan, dtype=np.float64)
    for target_index, target in enumerate(target_chainages):
        right = int(np.searchsorted(source_chainages, target, side="left"))
        if right < source_chainages.size and np.isclose(source_chainages[right], target, atol=1e-9, rtol=0.0):
            result[target_index] = source_values[right]
            continue
        if right == 0 or right >= source_chainages.size:
            continue
        left = right - 1
        left_value = source_values[left]
        right_value = source_values[right]
        if not np.isfinite(left_value) or not np.isfinite(right_value):
            continue
        fraction = (target - source_chainages[left]) / (source_chainages[right] - source_chainages[left])
        result[target_index] = left_value + fraction * (right_value - left_value)
    return result


def _sample_profile_table(
    *,
    line: LineString,
    order: int,
    scenario: str,
    profile_name: str,
    terrain_raster: Path,
    water_rasters: Mapping[str, Path],
    sampling_spacing: float,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Return a 50 m table and worst-1%-depth decision-summary record."""
    source_chainages, terrain = sample_raster_along_line(
        line,
        terrain_raster,
        sampling_spacing,
        method=SAMPLING_METHOD,
    )
    target_chainages = _exact_chainages(float(line.length))
    target_terrain = _interpolate_at_chainages(source_chainages, terrain, target_chainages)
    points = tuple(line.interpolate(float(chainage)) for chainage in target_chainages)

    data: dict[str, Any] = {
        "Order": np.full(target_chainages.shape, order, dtype=np.int64),
        "Scenario": np.full(target_chainages.shape, scenario, dtype=object),
        "Profile": np.full(target_chainages.shape, profile_name, dtype=object),
        "Chainage (m)": target_chainages,
        "Easting (m)": np.asarray([point.x for point in points], dtype=np.float64),
        "Northing (m)": np.asarray([point.y for point in points], dtype=np.float64),
        "Ground Level (mAHD)": target_terrain,
    }

    dense_water_by_aep: dict[str, FloatArray] = {}
    for aep, raster_path in water_rasters.items():
        water_chainages, water = sample_raster_along_line(
            line,
            raster_path,
            sampling_spacing,
            method=SAMPLING_METHOD,
        )
        if not np.array_equal(source_chainages, water_chainages):
            raise RuntimeError(f"Sampling stations differ between terrain and {aep} raster for {profile_name}.")
        water = interpolate_short_nan_gaps(water, max_gap=MAX_INTERPOLATION_GAP)
        dry = np.isnan(water) | np.isnan(terrain) | (water <= terrain)
        if DRY_AREA_HANDLING == "ground_level":
            water[dry] = terrain[dry]
        else:
            water[dry] = np.nan
        dense_water_by_aep[aep] = water
        target_water = _interpolate_at_chainages(source_chainages, water, target_chainages)
        data[_aep_column_name(aep)] = target_water
        data[_depth_column_name(aep)] = target_water - target_terrain

    reporting_water = dense_water_by_aep.get(REPORTING_AEP)
    if reporting_water is None:
        raise ValueError(f"Reporting AEP {REPORTING_AEP!r} is not among the sampled rasters {tuple(water_rasters)}.")
    reporting_depth = reporting_water - terrain
    valid_reporting = np.isfinite(reporting_depth)
    summary: dict[str, Any] = {
        "Order": order,
        "Scenario": scenario,
        "Profile": profile_name,
        "Reporting Basis": "Maximum 1% AEP depth along profile",
        "Sampling Spacing (m)": sampling_spacing,
    }
    if not np.any(valid_reporting):
        summary.update(
            {
                "Reporting Status": "No wet 1% AEP sample",
                "Chainage (m)": np.nan,
                "Easting (m)": np.nan,
                "Northing (m)": np.nan,
                "Ground Elevation (mAHD)": np.nan,
            }
        )
        reporting_index: int | None = None
    else:
        reporting_index = int(np.nanargmax(reporting_depth))
        reporting_chainage = float(source_chainages[reporting_index])
        reporting_point = line.interpolate(reporting_chainage)
        summary.update(
            {
                "Reporting Status": "Wet",
                "Chainage (m)": reporting_chainage,
                "Easting (m)": reporting_point.x,
                "Northing (m)": reporting_point.y,
                "Ground Elevation (mAHD)": terrain[reporting_index],
            }
        )
    for aep in SUMMARY_AEPS:
        event_label = _event_label(aep)
        event_water = dense_water_by_aep[aep]
        if reporting_index is None:
            water_level = np.nan
            depth = np.nan
        else:
            water_level = event_water[reporting_index]
            depth = water_level - terrain[reporting_index]
        summary[f"{event_label} Water Level (mAHD)"] = water_level
        summary[f"{event_label} Depth (m)"] = depth

    return pd.DataFrame(data), summary


def _collect_profile_tables() -> tuple[tuple[tuple[int, str, pd.DataFrame], ...], pd.DataFrame]:
    """Build and validate all configured profile tables and decision summaries."""
    terrain_crs, sampling_spacing = profiles._raster_crs_and_spacing(  # pyright: ignore[reportPrivateUsage]
        TERRAIN_TIF, SPACING
    )
    collected: dict[str, tuple[int, pd.DataFrame]] = {}
    summaries: list[dict[str, Any]] = []

    for scenario, (lines_gpkg, results_directory) in SCENARIOS.items():
        logger.info("Building 50 m profile tables for scenario {}", scenario)
        water_rasters = profiles.discover_tuflow_profile_rasters(
            results_directory,
            target_aeps=TARGET_AEPS,
            target_result_type=TARGET_RESULT_TYPE,
        )
        water_crs = profiles._read_water_raster_crs(water_rasters)  # pyright: ignore[reportPrivateUsage]
        raster_crs = profiles._resolve_known_raster_crs(terrain_crs, water_crs)  # pyright: ignore[reportPrivateUsage]
        lines, resolved_crs = profiles.load_profile_lines(
            lines_gpkg,
            requested_layer=LINES_LAYER_NAME,
            name_field=NAME_FIELD,
            lines_crs_if_missing=LINES_CRS_IF_MISSING,
            target_crs=raster_crs,
        )
        profiles._validate_known_profile_crs(resolved_crs)  # pyright: ignore[reportPrivateUsage]

        for _, row in lines.iterrows():
            raw_name = row[NAME_FIELD]
            if raw_name is None or bool(pd.isna(raw_name)) or not str(raw_name).strip():
                raise ValueError(f"Profile has an empty {NAME_FIELD!r} value.")
            line_name = str(raw_name)
            geometry = cast(BaseGeometry | None, row.geometry)
            if geometry is None:
                raise ValueError(f"Profile {line_name!r} has null geometry.")
            parts = profiles.split_profile_line(
                geometry,
                line_name=line_name,
                disconnected_handling=DISCONNECTED_LINE_HANDLING,
            )
            for part_number, line in enumerate(parts, start=1):
                display_line_name = line_name if len(parts) == 1 else f"{line_name} (part {part_number})"
                display_name = f"{scenario} - {display_line_name}"
                order = len(collected) + 1
                if display_name in collected:
                    raise ValueError(f"Duplicate profile table name: {display_name!r}")
                table, summary = _sample_profile_table(
                    line=line,
                    order=order,
                    scenario=scenario,
                    profile_name=display_line_name,
                    terrain_raster=TERRAIN_TIF,
                    water_rasters=water_rasters,
                    sampling_spacing=sampling_spacing,
                )
                collected[display_name] = (order, table)
                summaries.append(summary)

    ordered_tables = tuple(
        (order, display_name, table)
        for display_name, (order, table) in sorted(collected.items(), key=lambda item: item[1][0])
    )
    summary_table = pd.DataFrame(summaries).sort_values("Order").reset_index(drop=True)
    return ordered_tables, summary_table


def _pivot_table_data(
    tables: tuple[tuple[int, str, pd.DataFrame], ...],
) -> pd.DataFrame:
    """Return tidy one-row-per-profile-chainage-AEP data for PivotTables."""
    pivot_parts: list[pd.DataFrame] = []
    identity_columns = [
        "Order",
        "Scenario",
        "Profile",
        "Chainage (m)",
        "Easting (m)",
        "Northing (m)",
        "Ground Level (mAHD)",
    ]
    for _, _, table in tables:
        for aep in TARGET_AEPS:
            part = table.loc[:, identity_columns].copy()
            part["AEP"] = _aep_column_name(aep).removesuffix(" WL (mAHD)")
            part["Water Level (mAHD)"] = table[_aep_column_name(aep)]
            part["Depth (m)"] = table[_depth_column_name(aep)]
            pivot_parts.append(part)
    return pd.concat(pivot_parts, ignore_index=True)


def _summary_pivot_data(summary_table: pd.DataFrame) -> pd.DataFrame:
    """Return one row per reporting profile and AEP for summary PivotTables."""
    identity_columns = [
        "Order",
        "Scenario",
        "Profile",
        "Reporting Basis",
        "Reporting Status",
        "Sampling Spacing (m)",
        "Chainage (m)",
        "Easting (m)",
        "Northing (m)",
        "Ground Elevation (mAHD)",
    ]
    parts: list[pd.DataFrame] = []
    for aep in SUMMARY_AEPS:
        numeric = float(aep.casefold().removesuffix("p"))
        event_label = _event_label(aep)
        part = summary_table.loc[:, identity_columns].copy()
        part["AEP"] = event_label
        part["AEP (%)"] = numeric
        part["ARI Equivalent (years)"] = 100.0 / numeric
        part["Water Level (mAHD)"] = summary_table[f"{event_label} Water Level (mAHD)"]
        part["Depth (m)"] = summary_table[f"{event_label} Depth (m)"]
        parts.append(part)
    return pd.concat(parts, ignore_index=True)


def _format_worksheet(
    worksheet: Any,
    dataframe: pd.DataFrame,
    *,
    workbook: Any,
) -> None:
    """Apply compatible filters, widths, navigation, and number formats."""
    worksheet.freeze_panes(1, 0)
    worksheet.hide_gridlines(2)
    integer_format = workbook.add_format({"num_format": "0"})
    decimal_format = workbook.add_format({"num_format": f"0.{DECIMAL_PLACES * '0'}"})
    for column_index, header in enumerate(dataframe.columns):
        values = dataframe[header].dropna().astype(str)
        content_width = int(values.str.len().max()) if not values.empty else 0
        width = min(28, max(12, len(str(header)) + 2, content_width + 2))
        if header in {
            "Scenario",
            "Scenario",
            "Profile",
            "AEP",
            "Reporting Basis",
            "Reporting Status",
        }:
            cell_format = None
        elif header in {"Order", "Rows"}:
            cell_format = integer_format
        else:
            cell_format = decimal_format
        worksheet.set_column(column_index, column_index, width, cell_format)
    if not dataframe.empty:
        worksheet.autofilter(
            0,
            0,
            len(dataframe.index),
            len(dataframe.columns) - 1,
        )


def _write_outputs(
    tables: tuple[tuple[int, str, pd.DataFrame], ...],
    summary_table: pd.DataFrame,
    *,
    output_directory: Path,
) -> tuple[Path, ...]:
    """Write through same-folder staging files so Windows ACLs inherit correctly."""
    output_directory.parent.mkdir(parents=True, exist_ok=True)
    output_directory.mkdir(parents=True, exist_ok=True)
    completed: list[Path] = []
    staged_outputs: list[tuple[Path, Path]] = []
    try:
        for order, display_name, table in tables:
            safe_name = sanitize_windows_filename(display_name, fallback="profile")
            filename = f"{order:02d} - {safe_name}_50m.csv"
            final_csv = output_directory / filename
            temporary_csv = final_csv.with_name(f".{final_csv.stem}.tmp{final_csv.suffix}")
            staged_outputs.append((temporary_csv, final_csv))
            table.to_csv(
                temporary_csv,
                index=False,
                float_format=f"%.{DECIMAL_PLACES}f",
            )

        summary_pivot = _summary_pivot_data(summary_table)
        for filename, output_table in (
            (SUMMARY_CSV_NAME, summary_table),
            (SUMMARY_PIVOT_CSV_NAME, summary_pivot),
        ):
            final_csv = output_directory / filename
            temporary_csv = final_csv.with_name(f".{final_csv.stem}.tmp{final_csv.suffix}")
            staged_outputs.append((temporary_csv, final_csv))
            output_table.to_csv(
                temporary_csv,
                index=False,
                float_format=f"%.{DECIMAL_PLACES}f",
            )

        final_workbook_path = output_directory / WORKBOOK_NAME
        workbook_path = final_workbook_path.with_name(f".{final_workbook_path.stem}.tmp{final_workbook_path.suffix}")
        staged_outputs.append((workbook_path, final_workbook_path))
        all_profiles = pd.concat((table for _, _, table in tables), ignore_index=True)
        pivot_data = _pivot_table_data(tables)
        index_table = pd.DataFrame(
            {
                "Order": [order for order, _, _ in tables],
                "Profile": [display_name for _, display_name, _ in tables],
                "Rows": [len(table.index) for _, _, table in tables],
                "Interval (m)": SECTION_INTERVAL_M,
            }
        )
        with pd.ExcelWriter(workbook_path, engine="xlsxwriter") as writer:
            workbook = writer.book
            index_table.to_excel(writer, sheet_name="Index", index=False)
            _format_worksheet(
                writer.sheets["Index"],
                index_table,
                workbook=workbook,
            )
            summary_table.to_excel(writer, sheet_name="Decision Summary", index=False)
            _format_worksheet(
                writer.sheets["Decision Summary"],
                summary_table,
                workbook=workbook,
            )
            summary_pivot.to_excel(writer, sheet_name="Summary Pivot", index=False)
            _format_worksheet(
                writer.sheets["Summary Pivot"],
                summary_pivot,
                workbook=workbook,
            )
            pivot_data.to_excel(writer, sheet_name="Pivot Data", index=False)
            _format_worksheet(
                writer.sheets["Pivot Data"],
                pivot_data,
                workbook=workbook,
            )
            all_profiles.to_excel(writer, sheet_name="All Profiles", index=False)
            _format_worksheet(
                writer.sheets["All Profiles"],
                all_profiles,
                workbook=workbook,
            )
            for order, display_name, table in tables:
                sheet_name = f"{order:02d} {display_name}"[:31]
                table.to_excel(writer, sheet_name=sheet_name, index=False)
                _format_worksheet(
                    writer.sheets[sheet_name],
                    table,
                    workbook=workbook,
                )

        final_word_path = output_directory / WORD_DOCUMENT_NAME
        word_path = final_word_path.with_name(f".{final_word_path.stem}.tmp{final_word_path.suffix}")
        staged_outputs.append((word_path, final_word_path))
        _write_word_tables(word_path, tables)

        for temporary_path, final_path in staged_outputs:
            temporary_path.replace(final_path)
            completed.append(final_path)
    finally:
        for temporary_path, _ in staged_outputs:
            temporary_path.unlink(missing_ok=True)
    return tuple(completed)


def main() -> int:
    """Generate all configured 50 m profile tables."""
    configure_serial_logging(console_log_level="INFO")
    print(f"Water-level profile table exporter {SCRIPT_VERSION}")
    try:
        tables, summary_table = _collect_profile_tables()
        output_paths = _write_outputs(tables, summary_table, output_directory=OUTPUT_DIRECTORY)
    except (
        FileNotFoundError,
        FileExistsError,
        RasterSamplingError,
        RuntimeError,
        ValueError,
    ) as exc:
        logger.error("Profile table export failed: {}", exc)
        return 1
    except Exception:
        logger.exception("Profile table export failed")
        return 1

    row_count = sum(len(table.index) for _, _, table in tables)
    logger.success(
        "Created {} profile tables with {} total 50 m rows and {} output files in {}",
        len(tables),
        row_count,
        len(output_paths),
        OUTPUT_DIRECTORY,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
