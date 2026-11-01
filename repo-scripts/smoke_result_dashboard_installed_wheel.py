"""Smoke-test six copied TUFLOW result wrappers against an isolated installed wheel.

Run from package CI after the wheel has been installed under --expected-root.
All data and outputs are synthetic and disposable; no solver or submodule is used.
"""

from __future__ import annotations

import argparse
import os
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

import ryan_library
from ryan_library.mcp.installed_scripts import copy_wrapper


@dataclass(frozen=True, slots=True)
class Case:
    wrapper: str
    data_type: str
    sheet: str


CASES = (
    Case("TUFLOW-python/culvert_results/combine_culvert_maximums.py", "Cmx", "Maximums"),
    Case("TUFLOW-python/culvert_results/combine_culvert_timeseries.py", "Q", "1d_timeseries_data"),
    Case("TUFLOW-python/culvert_results/create_mean_culvert_peak_report.py", "Cmx", "aep-dur-mean"),
    Case("TUFLOW-python/culvert_results/create_median_culvert_peak_report.py", "Cmx", "aep-dur-med"),
    Case("TUFLOW-python/po_and_timeseries/combine_po_results.py", "PO", "combined_PO"),
    Case("TUFLOW-python/po_and_timeseries/combine_pomm_results.py", "POMM", "combined_POMM"),
)


def _write_inputs(folder: Path) -> None:
    folder.mkdir(parents=True)
    for tp, multiplier in ((1, 1.0), (2, 2.0)):
        run = f"Synthetic_01.00p_00120m_TP{tp:02d}"
        data = {
            "_1d_Cmx.csv": (
                "Chan ID,Time Qmax,Qmax,Vmax,Time Vmax\n"
                f"C01,0.5,{5 * multiplier},1.5,0.5\n"
            ),
            "_1d_Q.csv": (
                "Descriptor,Time,Q C01\n"
                f"row,0,{2 * multiplier}\n"
                f"row,1,{4 * multiplier}\n"
            ),
            "_PO.csv": (
                "Descriptor,Time,Flow\n"
                "Location,Time,C01\n"
                f"row,0,{2 * multiplier}\n"
                f"row,1,{4 * multiplier}\n"
            ),
            "_POMM.csv": (
                f"{run},Location,Flow\n"
                "value,Time,C01\n"
                f"value,Maximum (Extracted from Time Series),{5 * multiplier}\n"
                "value,Time of Maximum,0.5\n"
                "value,Minimum (Extracted From Time Series),-1\n"
                "value,Time of Minimum,0.0\n"
            ),
        }
        for suffix, content in data.items():
            (folder / f"{run}{suffix}").write_text(content, encoding="utf-8")


def _run(script: Path, work: Path, root: Path, *flags: str) -> subprocess.CompletedProcess[str]:
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(root)
    environment["PYTHONNOUSERSITE"] = "1"
    environment.pop("RYAN_TOOLS_REPOSITORY_ROOT", None)
    return subprocess.run(
        [
            sys.executable,
            str(script),
            "--working-directory",
            str(work),
            "--no-pause",
            "--paths-to-process",
            "inputs",
            *flags,
        ],
        cwd=work.parent,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
        stdin=subprocess.DEVNULL,
        check=False,
    )


def _snapshot(work: Path, sheet: str) -> tuple[tuple[str, ...], tuple[tuple[float, ...], ...]]:
    books = list(work.glob("*.xlsx"))
    if len(books) != 1:
        msg = f"Expected one workbook for {sheet}, found {books}"
        raise RuntimeError(msg)
    workbook = load_workbook(books[0], read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            msg = f"Missing output sheet {sheet}: {workbook.sheetnames}"
            raise RuntimeError(msg)
        rows = list(workbook[sheet].values)
        if len(rows) < 2:
            msg = f"No data rows in {sheet}: {rows}"
            raise RuntimeError(msg)
        headers = tuple(str(item) for item in rows[0])
        # Working-folder and path columns differ between the two disposable
        # copies. Compare generated numeric report values in row order.
        values = tuple(
            tuple(float(value) for value in row if isinstance(value, (int, float)) and not isinstance(value, bool))
            for row in rows[1:]
        )
        return headers, values
    finally:
        workbook.close()


def _one_case(case: Case, root: Path, work: Path, *, dashboard: bool) -> tuple[tuple[str, ...], tuple[tuple[float, ...], ...]]:
    work.mkdir(parents=True)
    _write_inputs(work / "inputs")
    script = copy_wrapper(case.wrapper, work)
    if not script.is_file() or not script.absolute().is_relative_to(work):
        msg = f"Wrapper was not copied into disposable work folder: {script}"
        raise RuntimeError(msg)
    help_result = _run(script, work, root, "--help")
    if help_result.returncode != 0 or "--live-dashboard" not in help_result.stdout:
        msg = f"Missing copied wrapper dashboard help for {case.wrapper}: {help_result.stderr[-900:]}"
        raise RuntimeError(msg)
    result = _run(
        script, work, root, "--data-types", case.data_type,
        "--live-dashboard" if dashboard else "--no-live-dashboard",
        "--log-file", "detail.log", "--file-log-level", "DEBUG",
    )
    if result.returncode != 0:
        msg = (
            f"Copied {case.wrapper} (dashboard={dashboard}) exited {result.returncode}: "
            f"stdout={result.stdout[-2600:]!r}, stderr={result.stderr[-1800:]!r}"
        )
        raise RuntimeError(msg)
    if "\x1b[" in result.stdout:
        msg = f"Redirected dashboard contained ANSI updates in {case.wrapper}"
        raise RuntimeError(msg)
    if dashboard and ": SUCCESS" not in result.stdout:
        msg = f"Missing full-workflow success in {case.wrapper}: {result.stdout[-1500:]}"
        raise RuntimeError(msg)
    if not (work / "detail.log").is_file():
        msg = f"Detailed file sink missing in {case.wrapper}"
        raise RuntimeError(msg)
    output = _snapshot(work, sheet=case.sheet)
    invalid = _run(script, work, root, "--data-types", case.data_type, "--live-max-rows", "0")
    if invalid.returncode == 0 or "live_max_rows" not in (invalid.stdout + invalid.stderr):
        msg = f"Invalid display settings were silently accepted by {case.wrapper}"
        raise RuntimeError(msg)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--expected-root", type=Path, required=True)
    args = parser.parse_args()
    # Physical identity check is intentional; user-configured input paths
    # elsewhere retain their original mapped-drive/UNC spelling.
    root = args.expected_root.resolve()
    actual = Path(ryan_library.__file__).resolve()
    if not actual.is_relative_to(root):
        msg = f"Result smoke imported checkout code, not installed wheel: {actual}"
        raise RuntimeError(msg)
    with TemporaryDirectory(prefix="ryan-tools-result-dashboards-") as location:
        base = Path(location)
        for index, case in enumerate(CASES):
            without = _one_case(case, root, base / f"{index}_ordinary", dashboard=False)
            with_dashboard = _one_case(case, root, base / f"{index}_dashboard", dashboard=True)
            if without != with_dashboard:
                msg = f"Dashboard changed workbook values/headers for {case.wrapper}: {without!r} vs {with_dashboard!r}"
                raise RuntimeError(msg)

    print(f"Six copied result wrappers passed installed-wheel smoke at {actual}: on/off workbook parity and settings.")


if __name__ == "__main__":
    main()
